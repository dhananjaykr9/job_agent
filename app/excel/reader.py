"""
Excel Reader — reads the existing Excel tracker to understand
its structure, existing data, and column layout.
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..utils.logger import get_logger

logger = get_logger("job_agent.excel.reader")


class ExcelReader:
    """Reads the existing Excel job tracker."""

    def __init__(self, file_path: str):
        self.file_path = Path(file_path)
        self._workbook = None
        self._sheet: Optional[Worksheet] = None

    def load(self) -> bool:
        """Load the workbook. Returns False if file doesn't exist."""
        if not self.file_path.exists():
            logger.error(f"Excel file not found: {self.file_path}")
            return False

        try:
            self._workbook = load_workbook(str(self.file_path))
            self._sheet = self._workbook.active
            logger.info(
                f"Loaded Excel: [bold]{self.file_path.name}[/] "
                f"({self._sheet.max_row} rows, {self._sheet.max_column} cols)"
            )
            return True
        except Exception as e:
            logger.error(f"Failed to load Excel: {e}")
            return False

    @property
    def headers(self) -> list[str]:
        """Get the header row (first row) as a list of strings."""
        if not self._sheet:
            return []
        return [
            str(cell.value).strip() if cell.value else ""
            for cell in self._sheet[1]
        ]

    @property
    def row_count(self) -> int:
        """Number of data rows (excluding header)."""
        if not self._sheet:
            return 0
        return max(0, self._sheet.max_row - 1)

    def get_existing_companies(self) -> set[str]:
        """Get all company names already in the spreadsheet."""
        if not self._sheet:
            return set()

        headers = self.headers
        try:
            col_idx = headers.index("Company") + 1
        except ValueError:
            return set()

        companies = set()
        for row in range(2, self._sheet.max_row + 1):
            val = self._sheet.cell(row=row, column=col_idx).value
            if val:
                companies.add(str(val).strip().lower())
        return companies

    def get_existing_urls(self) -> set[str]:
        """Get all job URLs already in the spreadsheet."""
        if not self._sheet:
            return set()

        headers = self.headers
        try:
            col_idx = headers.index("Job URL") + 1
        except ValueError:
            return set()

        urls = set()
        for row in range(2, self._sheet.max_row + 1):
            cell = self._sheet.cell(row=row, column=col_idx)
            # Check both cell value and hyperlink
            if cell.value:
                urls.add(str(cell.value).strip())
            if cell.hyperlink:
                urls.add(str(cell.hyperlink.target).strip())
        return urls

    def get_existing_entries(self) -> list[dict]:
        """Get all existing entries as list of dicts."""
        if not self._sheet:
            return []

        headers = self.headers
        entries = []
        for row in range(2, self._sheet.max_row + 1):
            entry = {}
            for col, header in enumerate(headers, 1):
                if header:
                    entry[header] = self._sheet.cell(row=row, column=col).value
            entries.append(entry)
        return entries

    def close(self):
        if self._workbook:
            self._workbook.close()
