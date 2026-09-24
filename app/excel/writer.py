"""
Excel Writer — appends new job rows to the existing Excel tracker.

CRITICAL RULES:
1. Never overwrite existing rows.
2. Never modify user-owned columns (Applied, Status, etc.).
3. Preserve all existing formatting.
4. Create a backup before writing.
5. Handle the "Job URL" column as a hyperlink.
"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.hyperlink import Hyperlink

from ..schemas import JobCandidate
from .mapper import ExcelMapper
from ..utils.logger import get_logger

logger = get_logger("job_agent.excel.writer")


class ExcelWriter:
    """Appends new job records to the existing Excel tracker."""

    def __init__(
        self,
        file_path: str,
        mapper: ExcelMapper,
    ):
        self.file_path = Path(file_path)
        self.mapper = mapper

    def _backup(self):
        """Create a timestamped backup before any write operation."""
        if not self.file_path.exists():
            return

        backup_dir = self.file_path.parent / "backups"
        backup_dir.mkdir(exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = backup_dir / f"{self.file_path.stem}_{timestamp}.xlsx"
        shutil.copy2(self.file_path, backup_path)
        logger.info(f"Backup created: {backup_path.name}")

    def write_jobs(self, jobs: list[JobCandidate]) -> int:
        """
        Append a list of jobs as new rows in the Excel file.

        Args:
            jobs: List of validated, non-duplicate JobCandidate objects.

        Returns:
            Number of rows successfully written.
        """
        if not jobs:
            logger.info("No jobs to write.")
            return 0

        if not self.file_path.exists():
            logger.info("Excel file not found — creating new file.")
            self._create_new_file()

        # ── Backup ──────────────────────────────────────
        self._backup()

        # ── Open workbook ───────────────────────────────
        wb = load_workbook(str(self.file_path))
        ws: Worksheet = wb.active

        headers = [
            str(cell.value).strip() if cell.value else ""
            for cell in ws[1]
        ]

        written = 0
        for job in jobs:
            row_data = self.mapper.job_to_row(job)
            next_row = ws.max_row + 1

            for col_idx, header in enumerate(headers, 1):
                if not header:
                    continue

                value = row_data.get(header)
                if value is None:
                    continue  # User-owned or unmapped — don't touch

                cell = ws.cell(row=next_row, column=col_idx)

                # Special handling for Job URL — make it a clickable hyperlink
                if header in ("Job URL", "Link", "URL", "Apply Link") and value:
                    cell.value = "Link"
                    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=value)
                    cell.font = Font(color="0563C1", underline="single")
                else:
                    cell.value = value

                # Copy alignment from the row above for consistency
                above_cell = ws.cell(row=next_row - 1, column=col_idx)
                if above_cell.alignment:
                    cell.alignment = Alignment(
                        horizontal=above_cell.alignment.horizontal,
                        vertical=above_cell.alignment.vertical,
                        wrap_text=above_cell.alignment.wrap_text,
                    )

            written += 1
            logger.info(
                f"[success]Written[/] row {next_row}: "
                f"{job.company_name} — {job.job_title}"
            )

        # ── Save ────────────────────────────────────────
        wb.save(str(self.file_path))
        wb.close()

        logger.info(
            f"Excel updated: [bold]{written} new rows[/] added to "
            f"{self.file_path.name}"
        )
        return written

    def _create_new_file(self):
        """Create a new Excel file with headers from the mapper."""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Job Tracker"

        for col_idx, header in enumerate(self.mapper.excel_headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True)

        wb.save(str(self.file_path))
        wb.close()
        logger.info(f"Created new Excel file: {self.file_path.name}")
